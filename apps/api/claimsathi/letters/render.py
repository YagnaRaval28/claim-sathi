"""PDF rendering with WeasyPrint (HarfBuzz shaping, so Gujarati and Devanagari render correctly). Imported lazily.

Fonts: the container installs Noto fonts (Dockerfile). All text is HTML-escaped; letters never contain markup.
"""

from __future__ import annotations

from html import escape
from typing import Any

CSS = """
@page { size: A4; margin: 25mm; @bottom-center { content: "Page " counter(page) " of " counter(pages);
        font-size: 9pt; color: #5F6A67; } }
body { font-family: "Noto Serif", "Noto Serif Gujarati", "Noto Serif Devanagari", serif; font-size: 11.5pt;
       line-height: 1.6; color: #1C2321; }
.meta { margin: 0; }
.subject { font-weight: 600; margin: 14pt 0; }
h3 { font-size: 11.5pt; font-weight: 700; margin: 12pt 0 4pt; }
blockquote { margin: 8pt 0 8pt 12pt; padding-left: 8pt; border-left: 2pt solid #1E6B40; }
sup { font-size: 8pt; }
ol, ul { padding-left: 16pt; list-style: none; }
.footnotes { margin-top: 18pt; border-top: 0.5pt solid #D8D4CB; padding-top: 6pt; font-size: 9pt; }
.disclaimer { margin-top: 10pt; font-size: 8.5pt; color: #5F6A67; }
.english-copy { break-before: page; }
.english-copy h2 { font-size: 12pt; margin: 0 0 12pt; }
"""


def _body_html(doc: dict[str, Any]) -> str:
    parts: list[str] = []
    for b in doc["blocks"]:
        t = b["type"]
        if t == "meta":
            parts.append(f'<p class="meta">{escape(b["text"])}</p>')
        elif t == "subject":
            parts.append(f'<p class="subject">{escape(b["text"])}</p>')
        elif t == "heading":
            parts.append(f"<h3>{escape(b['text'])}</h3>")
        elif t == "quote":
            parts.append(f'<blockquote lang="en">{escape(b["text"])}<sup>[{int(b["footnote"])}]</sup></blockquote>')
        elif t == "list":
            parts.append("<ul>" + "".join(f"<li>{escape(i)}</li>" for i in b["items"]) + "</ul>")
        else:
            parts.append(f"<p>{escape(b['text'])}</p>")
    notes = "".join(
        f"<p>[{int(n['n'])}] {escape(n['source_title'])}, {escape(n.get('clause_ref') or '')}, version "
        f"{escape(n['version'])}{' (SYNTHETIC test source)' if n.get('synthetic') else ''}</p>"
        for n in doc.get("footnotes", [])
    )
    return (
        f"{''.join(parts)}<div class='footnotes'>{notes}</div>"
        f"<p class='disclaimer'>{escape(doc.get('disclaimer', ''))}</p>"
    )


def letter_html(doc: dict[str, Any], language: str = "en", english_title: str = "English copy") -> str:
    """The letter, then (FR-024) its English copy on a new page when the letter is in Gujarati or Hindi."""
    copy = doc.get("english")
    tail = (
        f"<section lang='en' class='english-copy'><h2>{escape(english_title)}</h2>{_body_html(copy)}</section>"
        if copy
        else ""
    )
    return (
        f"<!doctype html><html lang='{escape(language)}'><head><meta charset='utf-8'><style>{CSS}</style></head>"
        f"<body>{_body_html(doc)}{tail}</body></html>"
    )


class RendererUnavailable(RuntimeError):
    """WeasyPrint's native libraries (Pango/GObject) are missing, e.g. native Windows without GTK."""


def render_pdf(doc: dict[str, Any], language: str = "en", english_title: str = "English copy") -> bytes:
    try:
        from weasyprint import HTML
    except OSError as exc:  # raised when libgobject/libpango cannot be loaded
        raise RendererUnavailable(type(exc).__name__) from exc

    pdf: bytes = HTML(string=letter_html(doc, language, english_title)).write_pdf()
    return pdf


def render_docx(doc: dict[str, Any], english_title: str = "English copy") -> bytes:
    """FR-028 DOCX: A4, Noto Serif (Gujarati/Devanagari fall back to the Noto Serif families), citation markers
    in the text and a numbered Sources list (python-docx cannot write native Word footnotes)."""
    import io

    from docx import Document as DocxDocument
    from docx.enum.section import WD_ORIENT
    from docx.shared import Mm, Pt

    d = DocxDocument()
    section = d.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(25))
    style = d.styles["Normal"]
    style.font.name = "Noto Serif"
    style.font.size = Pt(11.5)
    _docx_body(d, doc)
    copy = doc.get("english")
    if copy:  # FR-024: the English copy follows on a new page
        from docx.enum.text import WD_BREAK

        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        d.add_paragraph().add_run(english_title).bold = True
        _docx_body(d, copy)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _docx_body(d: Any, doc: dict[str, Any]) -> None:
    from docx.shared import Mm, Pt

    for b in doc["blocks"]:
        t = b["type"]
        if t == "list":
            for item in b.get("items", []):
                d.add_paragraph(str(item))
        elif t == "heading":
            p = d.add_paragraph()
            p.add_run(str(b.get("text", ""))).bold = True
        elif t == "quote":
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Mm(8)
            p.add_run(str(b.get("text", "")))
            marker = p.add_run(f"[{int(b['footnote'])}]")
            marker.font.superscript = True
        else:
            p = d.add_paragraph()
            run = p.add_run(str(b.get("text", "")))
            run.bold = t == "subject"
    notes = doc.get("footnotes", [])
    if notes:
        p = d.add_paragraph()
        p.add_run("Sources").bold = True
        for n in notes:
            synthetic = " (SYNTHETIC test source)" if n.get("synthetic") else ""
            d.add_paragraph(
                f"[{int(n['n'])}] {n['source_title']}, {n.get('clause_ref') or ''}, version {n['version']}{synthetic}"
            )
    disclaimer = d.add_paragraph()
    run = disclaimer.add_run(str(doc.get("disclaimer", "")))
    run.font.size = Pt(8.5)
